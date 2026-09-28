import os
import re
from io import BytesIO
from PIL import Image as PILImage
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as OpenPyXlImage
from openpyxl.utils import get_column_letter

# Ensure scripts temp dir for thumbnails
temp_thumb_dir = os.path.join("scripts", ".thumbs_cache")
os.makedirs(temp_thumb_dir, exist_ok=True)

# Map folder names to project info
PROJECT_METADATA = {
    "fashion/atlasi": {
        "discipline": "Fashion Photography",
        "project": "Atlasi",
        "code_prefix": "SA-FASH-ATL",
        "year": "2024",
        "desc": "A studio fashion story built around hand-woven textiles, sculptural silhouettes, and quiet gesture."
    },
    "fashion/sepidar": {
        "discipline": "Fashion Photography",
        "project": "Sepidar",
        "code_prefix": "SA-FASH-SEP",
        "year": "2020",
        "desc": "A fashion series of presence and fabric in open air — light, landscape, and quiet posture."
    },
    "fashion/zeeen": {
        "discipline": "Fashion Photography",
        "project": "Zeeen",
        "code_prefix": "SA-FASH-ZEE",
        "year": "2019",
        "desc": "A fashion story rooted in place — heritage textiles and the architecture of Persian light."
    },
    "food/gastronomie": {
        "discipline": "Food Photography",
        "project": "Gastronomie",
        "code_prefix": "SA-FOOD-GAS",
        "year": "2024",
        "desc": "A fine-dining editorial exploring texture, shadow, and the ritual of the plate."
    },
    "food/cuisine": {
        "discipline": "Food Photography",
        "project": "Cuisine",
        "code_prefix": "SA-FOOD-CUI",
        "year": "2024",
        "desc": "A culinary series of plated moments — steam, glaze, and the quiet geometry of the table."
    },
    "food/tasting": {
        "discipline": "Food Photography",
        "project": "Tasting",
        "code_prefix": "SA-FOOD-TAS",
        "year": "2024",
        "desc": "A tasting-menu story of color, glaze, and close-cropped appetite."
    },
    "portrait/calligraphy": {
        "discipline": "Portrait Photography",
        "project": "Calligraphy",
        "code_prefix": "SA-PORT-CAL",
        "year": "2024",
        "desc": "A studio portrait series pairing quiet gesture with Persian calligraphy — script as halo, script as light."
    },
    "portrait/photos": {
        "discipline": "Portrait Photography",
        "project": "Photos",
        "code_prefix": "SA-PORT-PHO",
        "year": "2024",
        "desc": "A quiet studio portrait series — warm light, presence, and steady gaze."
    },
    "portrait/gaze": {
        "discipline": "Portrait Photography",
        "project": "Gaze",
        "code_prefix": "SA-PORT-GAZ",
        "year": "2024",
        "desc": "A single portrait held in soft light — presence without distraction."
    },
    "product/objects": {
        "discipline": "Product Photography",
        "project": "Objects",
        "code_prefix": "SA-PROD-OBJ",
        "year": "2024",
        "desc": "A product photography series of crafted objects — form, material, and careful studio light."
    },
}

GRAPHIC_METADATA = {
    "graphic-design/book-covers": {
        "discipline": "Book Covers (Graphic Design)",
        "project": "Book Covers",
        "code_prefix": "SA-DES-BOK",
        "year": "2025",
        "desc": "A series of book cover designs for Persian titles — typography-led covers."
    },
    "graphic-design/posters": {
        "discipline": "Posters (Graphic Design)",
        "project": "Posters",
        "code_prefix": "SA-DES-POS",
        "year": "2025",
        "desc": "Poster designs for events, culture, and campaigns."
    }
}

def extract_number(filename):
    match = re.search(r'(\d+)', filename)
    return int(match.group(1)) if match else 999

def build_inventory_sheet(ws, metadata_map, sheet_title):
    ws.title = sheet_title

    # Header styling
    header_fill = PatternFill(start_color="0F0F0F", end_color="0F0F0F", fill_type="solid")
    header_font = Font(name="Arial", size=11, bold=True, color="3FEBCC")
    thin_border = Border(
        left=Side(style='thin', color="CCCCCC"),
        right=Side(style='thin', color="CCCCCC"),
        top=Side(style='thin', color="CCCCCC"),
        bottom=Side(style='thin', color="CCCCCC")
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)

    headers = [
        ("A", "Photo Code (Generated)", 24),
        ("B", "Thumbnail (Preview)", 18),
        ("C", "Discipline", 22),
        ("D", "Project Name", 18),
        ("E", "Current Filename", 26),
        ("F", "Proposed New Filename", 28),
        ("G", "Client", 20),
        ("H", "Date / Year", 15),
        ("I", "Model(s)", 25),
        ("J", "Makeup Artist", 22),
        ("K", "Photography Assistant", 24),
        ("L", "Stylist / Costume", 20),
        ("M", "Location / Studio", 20),
        ("N", "SEO Keywords / Tags", 30),
        ("O", "Photo Story / Alt Text", 35),
        ("P", "Asset Relative Path", 40)
    ]

    for col_letter, title, width in headers:
        cell = ws[f"{col_letter}1"]
        cell.value = title
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border
        ws.column_dimensions[col_letter].width = width

    ws.row_dimensions[1].height = 32

    current_row = 2

    for folder_rel, meta in metadata_map.items():
        folder_path = os.path.join("src", "assets", folder_rel.replace("/", os.sep))
        if not os.path.exists(folder_path):
            continue

        files = [
            f for f in os.listdir(folder_path)
            if f.lower().endswith(('.jpg', '.jpeg', '.png', '.webp')) and not f.endswith('-thumb.webp')
        ]
        files.sort(key=extract_number)

        for idx, filename in enumerate(files, start=1):
            full_path = os.path.join(folder_path, filename)
            photo_code = f"{meta['code_prefix']}-{idx:02d}"
            ext = os.path.splitext(filename)[1].lower()
            proposed_filename = f"{photo_code.lower()}{ext}"

            # Make thumbnail
            thumb_filename = f"thumb_{photo_code}{ext}"
            thumb_path = os.path.join(temp_thumb_dir, thumb_filename)
            try:
                with PILImage.open(full_path) as img:
                    img.thumbnail((80, 80))
                    # Convert to RGB if RGBA/P for JPG save
                    if img.mode in ("RGBA", "P") and ext in (".jpg", ".jpeg"):
                        img = img.convert("RGB")
                    img.save(thumb_path)
                has_thumb = True
            except Exception as e:
                print(f"Error creating thumb for {full_path}: {e}")
                has_thumb = False

            # Populate row
            row_data = [
                ("A", photo_code, center_align, True),
                ("B", "", center_align, False), # image goes here
                ("C", meta["discipline"], center_align, False),
                ("D", meta["project"], center_align, False),
                ("E", filename, left_align, False),
                ("F", proposed_filename, left_align, False),
                ("G", "", left_align, False), # Client (to fill)
                ("H", meta["year"], center_align, False), # Date/Year
                ("I", "", left_align, False), # Model (to fill)
                ("J", "", left_align, False), # Makeup artist (to fill)
                ("K", "", left_align, False), # Photography assistant (to fill)
                ("L", "", left_align, False), # Stylist (to fill)
                ("M", "Studio / Esfahan", left_align, False), # Location
                ("N", f"{meta['project']}, {meta['discipline']}, Siavash Akbari", left_align, False), # SEO Tags
                ("O", meta["desc"], left_align, False), # Description / Alt
                ("P", full_path.replace("\\", "/"), left_align, False)
            ]

            ws.row_dimensions[current_row].height = 70

            for col_letter, val, align, is_code in row_data:
                cell = ws[f"{col_letter}{current_row}"]
                cell.value = val
                cell.alignment = align
                cell.border = thin_border
                if is_code:
                    cell.font = Font(name="Arial", size=10, bold=True, color="008080")
                else:
                    cell.font = Font(name="Arial", size=10)

                # Light stripe for rows
                if current_row % 2 == 0:
                    cell.fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")

            # Insert image
            if has_thumb and os.path.exists(thumb_path):
                try:
                    img_obj = OpenPyXlImage(thumb_path)
                    img_obj.width = 65
                    img_obj.height = 65
                    ws.add_image(img_obj, f"B{current_row}")
                except Exception as e:
                    print(f"Error adding image to cell B{current_row}: {e}")

            current_row += 1

    # Freeze header row
    ws.freeze_panes = "A2"

def main():
    wb = openpyxl.Workbook()
    # Sheet 1: Photography
    ws_photo = wb.active
    build_inventory_sheet(ws_photo, PROJECT_METADATA, "Photography Portfolio")

    # Sheet 2: Graphic Design / Posters / Covers
    ws_graphic = wb.create_sheet()
    build_inventory_sheet(ws_graphic, GRAPHIC_METADATA, "Graphic Design & Posters")

    out_file = "PHOTO_INVENTORY_SIAVASH_AKBARI.xlsx"
    wb.save(out_file)
    print(f"Excel workbook successfully created: {out_file}")

if __name__ == "__main__":
    main()
